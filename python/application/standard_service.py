"""Đọc, kiểm tra và tra cứu tiêu chuẩn sản phẩm từ Excel."""

import math
import re
import sys
from pathlib import Path
from typing import Dict, Optional

from openpyxl import load_workbook

from domain.entities import StandardRecord

DEFAULT_STANDARD_FILENAME = "Upload%20standard.xlsx"
_HEADERS = {
    "part_no": "Mã hàng",
    "breakaway_max": "Tiêu chuẩn momen phá vỡ lớn nhất (Nm)",
    "operating_min": "Tiêu chuẩn momen hoạt động nhỏ nhất (Nm)",
    "operating_max": "Tiêu chuẩn momen hoạt động lớn nhất (Nm)",
    "internal_operating_min": "Tiêu chuẩn momen hoạt động nội bộ nhỏ nhất (Nm)",
    "internal_operating_max": "Tiêu chuẩn momen hoạt động nội bộ lớn nhất (Nm)",
    "lower_fixture": "Tiêu chuẩn đồ gá dưới",
    "upper_fixture": "Tiêu chuẩn đồ gá trên",
    "thread_code": "Mã ren",
    "special_warning": "Cảnh báo đặc biệt",
}


def _normalize_header(value: object) -> str:
    return re.sub(r"\s+", " ", str(value or "").strip()).casefold()


def default_standard_path() -> Path:
    """Trả đường dẫn file mặc định cho source hoặc bundle PyInstaller."""
    if getattr(sys, "frozen", False):
        base = Path(getattr(sys, "_MEIPASS", Path(sys.executable).parent))
    else:
        base = Path(__file__).resolve().parents[2]
    return base / "file" / DEFAULT_STANDARD_FILENAME


class StandardFileError(ValueError):
    """File tiêu chuẩn không hợp lệ; message chứa vị trí lỗi."""


class StandardService:
    """Cache tiêu chuẩn; chỉ thay cache sau khi file mới hợp lệ hoàn toàn."""

    def __init__(self, default_path: Optional[Path] = None):
        self._default_path = Path(default_path) if default_path else default_standard_path()
        self._active_path: Optional[Path] = None
        self._records: Dict[str, StandardRecord] = {}

    @property
    def default_path(self) -> Path:
        return self._default_path

    @property
    def active_path(self) -> Optional[Path]:
        return self._active_path

    @property
    def part_numbers(self) -> list[str]:
        """Danh sách các mã hàng có trong file tiêu chuẩn, đã sắp xếp."""
        return sorted(self._records.keys())

    def load(self, path: Optional[Path] = None) -> int:
        candidate = Path(path) if path else self._default_path
        records = self._parse(candidate)
        self._records = records
        self._active_path = candidate.resolve()
        return len(records)

    def lookup(self, part_no: str) -> Optional[StandardRecord]:
        key = str(part_no or "").strip().upper()
        return self._records.get(key) if len(key) == 8 else None

    def _parse(self, path: Path) -> Dict[str, StandardRecord]:
        if path.suffix.casefold() != ".xlsx":
            raise StandardFileError("File tiêu chuẩn phải có định dạng .xlsx")
        if not path.is_file():
            raise StandardFileError(f"Không tìm thấy file tiêu chuẩn: {path}")
        try:
            workbook = load_workbook(path, read_only=True, data_only=True)
        except Exception as exc:
            raise StandardFileError(f"Không đọc được file tiêu chuẩn: {exc}") from exc

        try:
            worksheet = next((s for s in workbook.worksheets if s.max_row > 0 and s.max_column > 0), None)
            if worksheet is None:
                raise StandardFileError("File tiêu chuẩn không có sheet dữ liệu")
            rows = worksheet.iter_rows(values_only=True)
            header_row = next(rows, None)
            if header_row is None:
                raise StandardFileError(f"Sheet '{worksheet.title}' không có tiêu đề")

            positions: Dict[str, int] = {}
            seen = set()
            expected = {_normalize_header(label): key for key, label in _HEADERS.items()}
            for index, value in enumerate(header_row):
                normalized = _normalize_header(value)
                if not normalized:
                    continue
                if normalized in seen:
                    raise StandardFileError(f"Sheet '{worksheet.title}', dòng 1: trùng tiêu đề '{value}'")
                seen.add(normalized)
                if normalized in expected:
                    positions[expected[normalized]] = index
            missing = [label for key, label in _HEADERS.items() if key not in positions]
            if missing:
                raise StandardFileError(f"Sheet '{worksheet.title}', dòng 1: thiếu cột {', '.join(missing)}")

            records: Dict[str, StandardRecord] = {}
            for row_number, row in enumerate(rows, start=2):
                if all(value in (None, "") for value in row):
                    continue

                def cell(key: str):
                    index = positions[key]
                    return row[index] if index < len(row) else None

                part_no = str(cell("part_no") or "").strip().upper()
                if len(part_no) != 8:
                    raise StandardFileError(f"Sheet '{worksheet.title}', dòng {row_number}, cột '{_HEADERS['part_no']}': mã hàng phải đúng 8 ký tự")
                if part_no in records:
                    raise StandardFileError(f"Sheet '{worksheet.title}', dòng {row_number}, cột '{_HEADERS['part_no']}': trùng mã '{part_no}'")

                numeric = {}
                for key in ("breakaway_max", "operating_min", "operating_max"):
                    try:
                        number = float(cell(key))
                    except (TypeError, ValueError) as exc:
                        raise StandardFileError(f"Sheet '{worksheet.title}', dòng {row_number}, cột '{_HEADERS[key]}': giá trị phải là số") from exc
                    if not math.isfinite(number):
                        raise StandardFileError(f"Sheet '{worksheet.title}', dòng {row_number}, cột '{_HEADERS[key]}': giá trị phải là số hữu hạn")
                    numeric[key] = number
                if numeric["breakaway_max"] < 0:
                    raise StandardFileError(f"Sheet '{worksheet.title}', dòng {row_number}, cột '{_HEADERS['breakaway_max']}': giá trị phải >= 0")
                if numeric["operating_min"] > numeric["operating_max"]:
                    raise StandardFileError(f"Sheet '{worksheet.title}', dòng {row_number}: '{_HEADERS['operating_min']}' phải <= '{_HEADERS['operating_max']}'")

                def parse_optional_float(key: str) -> Optional[float]:
                    val = cell(key)
                    if val is None or str(val).strip() == "":
                        return None
                    try:
                        num = float(val)
                    except (TypeError, ValueError) as exc:
                        raise StandardFileError(f"Sheet '{worksheet.title}', dòng {row_number}, cột '{_HEADERS[key]}': giá trị phải là số") from exc
                    if not math.isfinite(num):
                        raise StandardFileError(f"Sheet '{worksheet.title}', dòng {row_number}, cột '{_HEADERS[key]}': giá trị phải là số hữu hạn")
                    return num

                int_min = parse_optional_float("internal_operating_min")
                int_max = parse_optional_float("internal_operating_max")
                if int_min is not None and int_max is not None and int_min > int_max:
                    raise StandardFileError(f"Sheet '{worksheet.title}', dòng {row_number}: '{_HEADERS['internal_operating_min']}' phải <= '{_HEADERS['internal_operating_max']}'")

                def text(key: str) -> str:
                    value = cell(key)
                    return "" if value is None else str(value).strip()

                records[part_no] = StandardRecord(
                    part_no, numeric["breakaway_max"], numeric["operating_min"], numeric["operating_max"],
                    int_min, int_max,
                    text("lower_fixture"), text("upper_fixture"), text("thread_code"), text("special_warning"),
                )
            if not records:
                raise StandardFileError(f"Sheet '{worksheet.title}' không có bản ghi tiêu chuẩn")
            return records
        finally:
            workbook.close()
