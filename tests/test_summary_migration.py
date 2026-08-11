import hashlib, sys, tempfile, unittest
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch
from openpyxl import Workbook, load_workbook
ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / 'draw_plot'))
from draw_plot import TorquePlotViewer
class V:
    def __init__(self, v): self.v=v
    def value(self): return self.v
class L:
    def __init__(self, v): self.v=v
    def text(self): return self.v
class SummaryMigrationTests(unittest.TestCase):
    def viewer(self):
        v=TorquePlotViewer.__new__(TorquePlotViewer)
        v.machine_name='ZE-SG3-01'
        v.spec_min_spin=V(1); v.spec_max_spin=V(5); v.internal_spec_min_spin=V(2); v.internal_spec_max_spin=V(4)
        v.avg_label=L('3'); v.min_label=L('2.5'); v.max_label=L('3.5'); v.judgment_label=L('OK'); v.internal_judgment_label=L('OK')
        return v
    def metadata(self):
        return SimpleNamespace(test_item='Operating Torque',part_name='ITR',part_no='ABCDEFGH',sample_no=1,test_purpose='First',team='QM',line_no='01',tester='T',remark='')
    def test_migration_preserves_history_and_extra_columns(self):
        with tempfile.TemporaryDirectory() as d:
            p=Path(d)/'summary.xlsx'; wb=Workbook(); ws=wb.active
            ws.append(['Saved At','Date','Part No','Max Nm','Judgment','Custom']); ws.append(['old','2026-01-01','OLDPART1',9,'NG','keep']); wb.save(p); wb.close()
            self.assertEqual(str(p), self.viewer()._append_summary_report(str(p),self.metadata(),'new.xlsx','raw.csv','ctr.xlsx'))
            self.assertTrue(Path(str(p)+'.bak').exists()); wb=load_workbook(p); ws=wb.active; h=[c.value for c in ws[1]]
            self.assertNotIn('Date',h)
            max_index=h.index('Max Nm')
            self.assertEqual(h[max_index+1:max_index+5],['Internal spec Min','Internal spec Max','Judgment (Drawing spec)','Judgment (Internal spec)'])
            self.assertEqual(ws.cell(3,h.index('Machine')+1).value,'ZE-SG3-01')
            self.assertEqual(ws.cell(2,h.index('Part No')+1).value,'OLDPART1'); self.assertEqual(ws.cell(2,h.index('Judgment (Drawing spec)')+1).value,'NG')
            self.assertIsNone(ws.cell(2,h.index('Judgment (Internal spec)')+1).value); self.assertEqual(ws.cell(2,h.index('Custom')+1).value,'keep'); wb.close()
    def test_backup_failure_does_not_replace_original(self):
        with tempfile.TemporaryDirectory() as d:
            p=Path(d)/'summary.xlsx'; wb=Workbook(); wb.active.append(['Saved At','Judgment']); wb.active.append(['old','OK']); wb.save(p); wb.close(); before=hashlib.sha256(p.read_bytes()).digest()
            real_copy=__import__('shutil').copy2
            def fail_backup(src,dst,*a,**k):
                if str(dst).endswith('.bak'): raise OSError('denied')
                return real_copy(src,dst,*a,**k)
            with patch('shutil.copy2',side_effect=fail_backup), patch('PyQt5.QtWidgets.QMessageBox.warning'):
                self.assertIsNone(self.viewer()._append_summary_report(str(p),self.metadata(),'new.xlsx','raw.csv','ctr.xlsx'))
            self.assertEqual(before,hashlib.sha256(p.read_bytes()).digest())
if __name__=='__main__': unittest.main()


